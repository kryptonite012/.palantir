"""Streaming parsers for tabular and semi-structured formats.

Every parser yields ``dict[str, Any]`` rows with string column names. Values
are left as close to the source as possible (strings for text formats);
typing happens in normalization so that original values are preserved.
"""

from __future__ import annotations

import csv
import io
import json
import sqlite3
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

from project_atlas.core.errors import ValidationFailed

TABULAR_FORMATS = {"csv", "tsv", "xlsx", "json", "jsonl", "xml", "sqlite", "parquet"}
DOCUMENT_FORMATS = {"txt", "md", "pdf", "docx", "eml"}
GEO_FORMATS = {"geojson"}
ALLOWED_EXTENSIONS = {
    ".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".xlsx": "xlsx", ".json": "json", ".jsonl": "jsonl",
    ".ndjson": "jsonl", ".xml": "xml", ".txt": "txt", ".md": "md", ".markdown": "md", ".pdf": "pdf",
    ".docx": "docx", ".eml": "eml", ".sqlite": "sqlite", ".sqlite3": "sqlite", ".db": "sqlite",
    ".parquet": "parquet", ".geojson": "geojson",
}

# Magic-byte checks: the declared extension must agree with the content.
_MAGIC = {
    "pdf": [b"%PDF-"], "xlsx": [b"PK\x03\x04"], "docx": [b"PK\x03\x04"], "sqlite": [b"SQLite format 3\x00"],
    "parquet": [b"PAR1"],
}

ENCODINGS = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]


def detect_format(filename: str, head: bytes) -> str:
    ext = Path(filename).suffix.lower()
    fmt = ALLOWED_EXTENSIONS.get(ext)
    if fmt is None:
        raise ValidationFailed(f"File type {ext or '(none)'} is not permitted")
    magics = _MAGIC.get(fmt)
    if magics and not any(head.startswith(m) for m in magics):
        raise ValidationFailed(f"File content does not match its {ext} extension")
    if fmt not in _MAGIC and b"\x00" in head[:4096] and fmt not in {"eml"}:
        # text formats should not contain NUL bytes (UTF-16 text is not supported)
        raise ValidationFailed("Text file contains binary content (NUL bytes); UTF-16 is not supported")
    return fmt


def detect_encoding(sample: bytes) -> tuple[str, bool]:
    """Return (encoding, had_errors). Tries strict decoders in order."""
    for enc in ENCODINGS:
        try:
            sample.decode(enc)
            return enc, False
        except UnicodeDecodeError as exc:
            # a multi-byte sequence may be cut at the sample boundary
            if enc.startswith("utf-8") and exc.start >= len(sample) - 3:
                return enc, False
            continue
    return "utf-8", True


@dataclass
class ParseOptions:
    encoding: str | None = None
    delimiter: str | None = None
    sheet: str | None = None
    table: str | None = None  # for sqlite sources
    record_path: str | None = None  # dotted path to the array in a JSON document
    xml_record_tag: str | None = None
    has_header: bool = True
    extra: dict[str, Any] = field(default_factory=dict)

    @classmethod
    def from_dict(cls, d: dict[str, Any] | None) -> "ParseOptions":
        d = dict(d or {})
        known = {k: d.pop(k) for k in list(d) if k in cls.__dataclass_fields__ and k != "extra"}
        return cls(**known, extra=d)


@dataclass
class Inspection:
    format: str
    encoding: str | None
    encoding_errors: bool
    delimiter: str | None
    columns: list[str]
    preview: list[dict[str, Any]]
    sheets: list[str] = field(default_factory=list)
    tables: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _sniff_delimiter(text: str, fmt: str) -> str:
    if fmt == "tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _text_stream(path: Path, encoding: str) -> io.TextIOWrapper:
    return open(path, "r", encoding=encoding, errors="replace", newline="")


def _dedupe_headers(headers: list[str]) -> list[str]:
    out, seen = [], {}
    for i, h in enumerate(headers):
        h = (h or "").strip() or f"column_{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 1
        out.append(h)
    return out


def _flatten(obj: Any, prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, dict):
                out.update(_flatten(v, key))
            elif isinstance(v, list):
                out[key] = json.dumps(v, ensure_ascii=False)
            else:
                out[key] = v
    else:
        out[prefix or "value"] = obj
    return out


def _resolve_path(doc: Any, path: str | None) -> Any:
    if not path:
        if isinstance(doc, list):
            return doc
        if isinstance(doc, dict):
            # Use the first list-valued key if the document wraps its records.
            for v in doc.values():
                if isinstance(v, list) and v and isinstance(v[0], dict):
                    return v
            return [doc]
        raise ValidationFailed("JSON document does not contain records")
    cur = doc
    for part in path.split("."):
        if not isinstance(cur, dict) or part not in cur:
            raise ValidationFailed(f"record_path {path!r} not found in JSON document")
        cur = cur[part]
    if not isinstance(cur, list):
        raise ValidationFailed(f"record_path {path!r} does not point to an array")
    return cur


def iter_rows(path: Path, fmt: str, opts: ParseOptions) -> Iterator[dict[str, Any]]:
    if fmt in {"csv", "tsv"}:
        enc = opts.encoding or detect_encoding(path.read_bytes()[:65536])[0]
        with _text_stream(path, enc) as f:
            sample = f.read(20000)
            f.seek(0)
            delim = opts.delimiter or _sniff_delimiter(sample, fmt)
            reader = csv.reader(f, delimiter=delim)
            headers: list[str] | None = None
            for row in reader:
                if headers is None:
                    if opts.has_header:
                        headers = _dedupe_headers(row)
                        continue
                    headers = [f"column_{i + 1}" for i in range(len(row))]
                if not any(c.strip() for c in row):
                    continue
                if len(row) < len(headers):
                    row = row + [""] * (len(headers) - len(row))
                rec = dict(zip(headers, row))
                if len(row) > len(headers):
                    rec["__extra_columns__"] = row[len(headers):]
                yield rec
    elif fmt == "jsonl":
        enc = opts.encoding or "utf-8"
        with _text_stream(path, enc) as f:
            for n, line in enumerate(f, 1):
                line = line.strip()
                if not line:
                    continue
                try:
                    yield _flatten(json.loads(line))
                except json.JSONDecodeError as exc:
                    yield {"__parse_error__": f"line {n}: {exc.msg}"}
    elif fmt == "json":
        enc = opts.encoding or "utf-8-sig"
        try:
            doc = json.loads(path.read_text(encoding=enc, errors="replace"))
        except json.JSONDecodeError as exc:
            raise ValidationFailed(f"Invalid JSON: {exc.msg} at line {exc.lineno}") from exc
        for item in _resolve_path(doc, opts.record_path):
            yield _flatten(item)
    elif fmt == "xml":
        yield from _iter_xml(path, opts)
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[opts.sheet] if opts.sheet else wb.worksheets[0]
            headers = None
            for row in ws.iter_rows(values_only=True):
                if headers is None:
                    headers = _dedupe_headers([str(c) if c is not None else "" for c in row])
                    continue
                if all(c is None or str(c).strip() == "" for c in row):
                    continue
                yield {h: ("" if v is None else v) for h, v in zip(headers, row)}
        finally:
            wb.close()
    elif fmt == "sqlite":
        conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        try:
            tables = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
            table = opts.table or (tables[0] if tables else None)
            if table not in tables:
                raise ValidationFailed(f"Table {table!r} not found; available: {tables}")
            cur = conn.execute(f'SELECT * FROM "{table}"')  # table name validated against catalog above
            cols = [d[0] for d in cur.description]
            for row in cur:
                yield dict(zip(cols, row))
        finally:
            conn.close()
    elif fmt == "parquet":
        import pyarrow.parquet as pq

        pf = pq.ParquetFile(path)
        for batch in pf.iter_batches(batch_size=2000):
            for row in batch.to_pylist():
                yield {str(k): v for k, v in row.items()}
    else:
        raise ValidationFailed(f"Format {fmt} is not a tabular format")


def _iter_xml(path: Path, opts: ParseOptions) -> Iterator[dict[str, Any]]:
    # Defend against entity-expansion attacks: refuse DTDs outright.
    head = path.read_bytes()[:4096].lower()
    if b"<!doctype" in head or b"<!entity" in head:
        raise ValidationFailed("XML documents with DTD/entity declarations are not accepted")
    tag = opts.xml_record_tag
    if tag is None:
        root = ET.parse(path).getroot()
        children = list(root)
        if not children:
            return
        tag = children[0].tag
    for _event, elem in ET.iterparse(path, events=("end",)):
        if elem.tag == tag:
            rec: dict[str, Any] = {f"@{k}": v for k, v in elem.attrib.items()}
            for child in elem:
                if len(child):
                    for sub in child:
                        rec[f"{child.tag}.{sub.tag}"] = (sub.text or "").strip()
                else:
                    rec[child.tag] = (child.text or "").strip()
            if not len(elem) and elem.text and elem.text.strip():
                rec["value"] = elem.text.strip()
            yield rec
            elem.clear()


def inspect(path: Path, fmt: str, opts: ParseOptions, preview_rows: int = 20) -> Inspection:
    notes: list[str] = []
    enc, enc_err, delim = None, False, None
    sheets: list[str] = []
    tables: list[str] = []
    if fmt in {"csv", "tsv", "jsonl", "json", "xml"}:
        enc, enc_err = detect_encoding(path.read_bytes()[:65536])
        if opts.encoding:
            enc = opts.encoding
        if enc_err:
            notes.append("Encoding could not be determined strictly; undecodable bytes will be replaced")
        opts.encoding = enc
    if fmt in {"csv", "tsv"}:
        with _text_stream(path, enc or "utf-8") as f:
            delim = opts.delimiter or _sniff_delimiter(f.read(20000), fmt)
        opts.delimiter = delim
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        sheets = wb.sheetnames
        wb.close()
    if fmt == "sqlite":
        conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        tables = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        conn.close()
    preview: list[dict[str, Any]] = []
    columns: list[str] = []
    for i, row in enumerate(iter_rows(path, fmt, opts)):
        if i >= preview_rows:
            break
        preview.append(row)
        for k in row:
            if k not in columns and not k.startswith("__"):
                columns.append(k)
    return Inspection(fmt, enc, enc_err, delim, columns, preview, sheets, tables, notes)


def jsonable(v: Any) -> Any:
    """Convert parser output to JSON-safe values while keeping their meaning."""
    from datetime import date, datetime, time
    from decimal import Decimal

    if v is None or isinstance(v, (str, int, float, bool)):
        if isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
            return None
        return v
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bytes):
        return v.hex()
    if isinstance(v, (list, dict)):
        return json.loads(json.dumps(v, default=str))
    return str(v)
