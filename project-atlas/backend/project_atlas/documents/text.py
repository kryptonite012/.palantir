"""Deterministic text utilities: extraction, chunking, language detection, keywords, candidates, similarity.

Methods here are heuristics and are labelled as such in the UI and API output.
"""

from __future__ import annotations

import email
import email.policy
import html
import math
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

STOPWORDS: dict[str, set[str]] = {
    "en": set("the and of to in is that for it as with was on be by this are or at from an not have has had which but were "
              "their they you we he she his her its been will would there all can if no more one about".split()),
    "es": set("el la de que y en los del se las por un para con una su al lo como más pero sus le ya o este sí porque "
              "esta entre cuando muy sin sobre también me hasta hay donde".split()),
    "fr": set("le la les de des et en un une du que qui dans pour pas sur au avec ce il elle ne se plus par sont est "
              "mais ou nous vous leur été être".split()),
    "de": set("der die das und in den von zu mit sich des auf für ist im dem nicht ein eine als auch es an werden aus "
              "er hat dass sie nach wird bei einer um am sind noch wie".split()),
    "it": set("il di che e la per un in è non una sono del della con si le da al lo ma come anche nel più gli".split()),
    "pt": set("o a de que e do da em um para é com não uma os no se na por mais as dos como mas foi ao ele das".split()),
    "nl": set("de het een en van in is dat op te zijn voor met die niet aan er maar om ook als bij of dan nog".split()),
}
ALL_STOP = set().union(*STOPWORDS.values())
WORD_RE = re.compile(r"[^\W\d_][\w'-]*", re.UNICODE)
INJECTION_RE = re.compile(
    r"(ignore (all |any )?(previous|prior|above) (instructions|prompts)|disregard (the )?(system|previous)|"
    r"you are now|system prompt|reveal (your|the) (instructions|prompt)|act as (an? )?(admin|administrator)|"
    r"execute (the following|this) (command|sql)|drop table|delete from)",
    re.IGNORECASE,
)

CLASSIFICATION_RULES = [
    ("invoice", ["invoice", "amount due", "bill to", "payment terms"]),
    ("contract", ["agreement", "hereinafter", "party", "terms and conditions", "governing law"]),
    ("correspondence", ["dear", "regards", "sincerely", "subject:"]),
    ("report", ["executive summary", "findings", "methodology", "conclusion"]),
    ("meeting_notes", ["minutes", "attendees", "action items", "agenda"]),
    ("financial_statement", ["balance sheet", "income statement", "cash flow", "statement of account"]),
]


def detect_language(text: str) -> tuple[str | None, float]:
    words = [w.lower() for w in WORD_RE.findall(text[:20000])]
    if len(words) < 20:
        return None, 0.0
    scores = {lang: sum(1 for w in words if w in sw) for lang, sw in STOPWORDS.items()}
    best = max(scores, key=scores.get)
    total = sum(scores.values())
    if scores[best] < 3 or total == 0:
        return None, 0.0
    return best, round(scores[best] / total, 3)


def keywords(text: str, k: int = 15) -> list[dict[str, Any]]:
    words = [w.lower() for w in WORD_RE.findall(text)]
    counts = Counter(w for w in words if len(w) >= 3 and w not in ALL_STOP)
    n = max(sum(counts.values()), 1)
    return [{"term": t, "count": c, "tf": round(c / n, 5)} for t, c in counts.most_common(k)]


def classify(text: str) -> tuple[str, list[str]]:
    low = text.lower()
    best, hits_best = "unclassified", []
    for label, cues in CLASSIFICATION_RULES:
        hits = [c for c in cues if c in low]
        if len(hits) >= 2 and len(hits) > len(hits_best):
            best, hits_best = label, hits
    return best, hits_best


_EMAIL = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
_PHONE = re.compile(r"(?<!\w)(\+?\d[\d\s().-]{7,}\d)(?!\w)")
_DATE = re.compile(r"\b(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2} (?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* \d{4})\b")
_MONEY = re.compile(r"(?:[$€£]\s?\d[\d,]*(?:\.\d+)?|\b\d[\d,]*(?:\.\d+)?\s?(?:USD|EUR|GBP|NOK|CHF)\b)")
_PROPER = re.compile(r"\b([A-Z][a-z]+(?:\s+(?:[A-Z][a-z]+|[A-Z]\.|of|de|van|von|and|&)){1,4}\s*(?:[A-Z][a-z]+|Ltd|Inc|LLC|GmbH|AG)?)\b")


def entity_candidates(pages: list[tuple[int, str]], gazetteer: dict[str, tuple[str, str]]) -> list[dict[str, Any]]:
    """Candidates only; never merged into the graph without review."""
    found: dict[tuple[str, str], dict[str, Any]] = {}

    def add(kind: str, value: str, page: int, method: str, entity_id: str | None = None):
        key = (kind, value.strip())
        if key not in found:
            found[key] = {"kind": kind, "text": value.strip(), "pages": [], "count": 0, "method": method,
                          "matched_entity_id": entity_id, "status": "candidate"}
        c = found[key]
        c["count"] += 1
        if page not in c["pages"]:
            c["pages"].append(page)

    for page, text in pages:
        for m in _EMAIL.finditer(text):
            add("email", m.group(0), page, "regex")
        for m in _PHONE.finditer(text):
            digits = re.sub(r"\D", "", m.group(1))
            if 8 <= len(digits) <= 15:
                add("phone", m.group(1), page, "regex")
        for m in _DATE.finditer(text):
            add("date", m.group(0), page, "regex")
        for m in _MONEY.finditer(text):
            add("money", m.group(0), page, "regex")
        for m in _PROPER.finditer(text):
            v = m.group(1).strip()
            if v.split()[0].lower() not in ALL_STOP and len(v) > 4:
                add("name", v, page, "capitalization heuristic")
        low = text.lower()
        for name_low, (eid, display) in gazetteer.items():
            if len(name_low) >= 4 and re.search(r"\b" + re.escape(name_low) + r"\b", low):
                add("known_entity", display, page, "exact match against project entities", eid)
    return sorted(found.values(), key=lambda c: (-c["count"], c["kind"], c["text"]))[:300]


def chunk_text(text: str, target: int = 900) -> list[tuple[int, str]]:
    """Split into chunks near paragraph/sentence boundaries; returns (char_start, text)."""
    out: list[tuple[int, str]] = []
    text = text.replace("\r\n", "\n")
    pos = 0
    n = len(text)
    while pos < n:
        end = min(n, pos + target)
        if end < n:
            window = text[pos:end + 200]
            cut = max(window.rfind("\n\n", 0, target + 150), window.rfind(". ", 0, target + 150))
            if cut > target // 3:
                end = pos + cut + 1
        piece = text[pos:end].strip()
        if piece:
            out.append((pos, piece))
        pos = end
    return out


def shingles(text: str, k: int = 5) -> set[int]:
    words = [w.lower() for w in WORD_RE.findall(text)]
    return {hash(" ".join(words[i:i + k])) for i in range(max(0, len(words) - k + 1))}


def jaccard(a: set[int], b: set[int]) -> float:
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def tfidf_vectors(texts: list[str]) -> list[dict[str, float]]:
    docs = [Counter(w.lower() for w in WORD_RE.findall(t) if w.lower() not in ALL_STOP and len(w) > 2) for t in texts]
    df = Counter(term for d in docs for term in d)
    n = len(docs)
    vecs = []
    for d in docs:
        v = {t: (c / max(sum(d.values()), 1)) * math.log((1 + n) / (1 + df[t]) + 1) for t, c in d.items()}
        norm = math.sqrt(sum(x * x for x in v.values())) or 1.0
        vecs.append({t: x / norm for t, x in v.items()})
    return vecs


def cosine(a: dict[str, float], b: dict[str, float]) -> float:
    if len(a) > len(b):
        a, b = b, a
    return sum(v * b.get(t, 0.0) for t, v in a.items())


# --------------------------------------------------------------------------- extraction


def ocr_available() -> bool:
    return bool(shutil.which("tesseract") and shutil.which("pdftoppm"))


def _ocr_pdf_page(path: Path, page: int) -> str:
    with tempfile.TemporaryDirectory() as td:
        base = Path(td) / "p"
        subprocess.run(["pdftoppm", "-r", "200", "-f", str(page), "-l", str(page), "-png", str(path), str(base)],
                       check=True, capture_output=True, timeout=120)
        imgs = sorted(Path(td).glob("p*.png"))
        if not imgs:
            return ""
        res = subprocess.run(["tesseract", str(imgs[0]), "-", "--psm", "3"], check=True, capture_output=True, timeout=300)
        return res.stdout.decode("utf-8", errors="replace")


def extract(path: Path, fmt: str) -> dict[str, Any]:
    """Return {pages: [(page_no, text)], metadata, ocr_status, page_label}."""
    meta: dict[str, Any] = {}
    ocr_status = "not_required"
    if fmt == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        pages = []
        needs_ocr = []
        for i, pg in enumerate(reader.pages, 1):
            t = pg.extract_text() or ""
            if len(t.strip()) < 20:
                needs_ocr.append(i)
            pages.append((i, t))
        if needs_ocr:
            if ocr_available():
                for i in needs_ocr:
                    pages[i - 1] = (i, _ocr_pdf_page(path, i))
                ocr_status = f"ocr_applied:{len(needs_ocr)}"
            else:
                ocr_status = f"ocr_required_unavailable:{len(needs_ocr)}"
        info = reader.metadata or {}
        meta = {"pdf_title": str(info.get("/Title", "") or ""), "pdf_author": str(info.get("/Author", "") or ""),
                "pages": len(reader.pages)}
        return {"pages": pages, "metadata": meta, "ocr_status": ocr_status, "page_label": "page"}
    if fmt == "docx":
        import docx

        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                paras.append(" | ".join(c.text.strip() for c in row.cells))
        sections, cur, size = [], [], 0
        for p in paras:
            cur.append(p)
            size += len(p)
            if size >= 3000:
                sections.append("\n\n".join(cur))
                cur, size = [], 0
        if cur:
            sections.append("\n\n".join(cur))
        cp = d.core_properties
        meta = {"docx_title": cp.title or "", "docx_author": cp.author or "", "paragraphs": len(paras)}
        return {"pages": [(i, s) for i, s in enumerate(sections, 1)], "metadata": meta, "ocr_status": ocr_status,
                "page_label": "section"}
    if fmt in {"txt", "md"}:
        raw = path.read_bytes()
        from project_atlas.ingestion.parsers import detect_encoding

        enc, _ = detect_encoding(raw[:65536])
        text = raw.decode(enc, errors="replace")
        # split long text files into ~3000-char sections at blank lines for page-level references
        parts, cur = [], ""
        for block in re.split(r"\n\s*\n", text):
            if len(cur) + len(block) > 3000 and cur:
                parts.append(cur)
                cur = ""
            cur += block + "\n\n"
        if cur.strip():
            parts.append(cur)
        return {"pages": [(i, p) for i, p in enumerate(parts, 1)], "metadata": {"encoding": enc},
                "ocr_status": ocr_status, "page_label": "section"}
    if fmt == "eml":
        msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
        meta = {k: str(msg.get(k, "")) for k in ("From", "To", "Cc", "Date", "Subject", "Message-ID")}
        body = msg.get_body(preferencelist=("plain", "html"))
        text = body.get_content() if body else ""
        if body is not None and body.get_content_type() == "text/html":
            text = html.unescape(re.sub(r"<[^>]+>", " ", text))
        attachments = [part.get_filename() for part in msg.iter_attachments()]
        meta["attachments"] = [a for a in attachments if a]
        header = "\n".join(f"{k}: {v}" for k, v in meta.items() if k != "attachments" and v)
        return {"pages": [(1, header + "\n\n" + text)], "metadata": meta, "ocr_status": ocr_status, "page_label": "message"}
    raise ValueError(f"Unsupported document format {fmt}")
