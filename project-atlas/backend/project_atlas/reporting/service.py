"""Report builder: resolves a report definition into content blocks and renders PDF/HTML/JSON/CSV/XLSX.

Every export writes a manifest (dataset versions, analysis runs and parameters,
timestamps, software versions, output SHA-256) alongside the file.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import platform
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from project_atlas import APP_NAME, __version__
from project_atlas.core.errors import NotFound, ValidationFailed
from project_atlas.database.models import (
    AnalysisRun, Dataset, DatasetVersion, Entity, Event, Evidence, Finding, FindingEvidence, GeoFeature, GraphSnapshot,
    Investigation, Record, Relationship, Report, ReportExport,
)

if TYPE_CHECKING:
    from project_atlas.core.context import AtlasContext

RENDERER_VERSION = "report-1.0"
FONT_DIR = Path(__file__).parent / "fonts"
FORMATS = {"pdf", "html", "json", "csv", "xlsx"}
SECTION_TYPES = {"summary", "methodology", "text", "limitations", "appendix", "findings", "analysis", "dataset_table",
                 "graph", "timeline", "map", "evidence"}


def validate_definition(defn: dict[str, Any]) -> None:
    secs = defn.get("sections")
    if not isinstance(secs, list) or not secs:
        raise ValidationFailed("Report needs at least one section")
    for i, s in enumerate(secs):
        if s.get("type") not in SECTION_TYPES:
            raise ValidationFailed(f"sections[{i}]: unknown type {s.get('type')!r}")


# --------------------------------------------------------------------------- resolution


class Resolver:
    def __init__(self, db: Session, project_id: str, perms: set[str]):
        self.db, self.project_id, self.perms = db, project_id, perms
        self.manifest: dict[str, Any] = {"dataset_versions": {}, "analysis_runs": {}, "graph_snapshots": [],
                                         "investigations": [], "redactions": []}

    def _dataset(self, dataset_id: str, version: int | None) -> tuple[Dataset, DatasetVersion]:
        ds = self.db.get(Dataset, dataset_id)
        if ds is None or ds.project_id != self.project_id or (
                ds.access_policy == "restricted" and "dataset.read_restricted" not in self.perms):
            raise NotFound("Dataset referenced by report not found or not accessible")
        v = version or ds.current_version
        dv = self.db.execute(select(DatasetVersion).where(DatasetVersion.dataset_id == ds.id, DatasetVersion.version == v)).scalar_one()
        self.manifest["dataset_versions"][dv.id] = {"dataset_id": ds.id, "name": ds.name, "version": dv.version,
                                                    "content_sha256": dv.content_sha256, "synthetic": ds.is_synthetic}
        return ds, dv

    def _hidden(self, dv: DatasetVersion) -> set[str]:
        if "sensitive.export" in self.perms:
            return set()
        hidden = {f["name"] for f in dv.schema_json.get("fields", []) if f.get("sensitive")}
        if hidden:
            self.manifest["redactions"].append({"dataset_version_id": dv.id, "fields": sorted(hidden)})
        return hidden

    def resolve(self, defn: dict[str, Any]) -> list[dict[str, Any]]:
        blocks: list[dict[str, Any]] = []
        for s in defn["sections"]:
            title = s.get("title") or s["type"].replace("_", " ").title()
            blocks.append({"kind": "heading", "text": title})
            fn = getattr(self, f"_s_{s['type']}", None)
            blocks.extend(fn(s) if fn else [{"kind": "paragraph", "text": s.get("body", "")}])
        return blocks

    def _text(self, s):
        return [{"kind": "paragraph", "text": p} for p in str(s.get("body", "")).split("\n\n") if p.strip()]

    _s_summary = _s_methodology = _s_text = _s_limitations = _s_appendix = _text

    def _s_dataset_table(self, s):
        ds, dv = self._dataset(s["dataset_id"], s.get("version"))
        hidden = self._hidden(dv)
        cols = s.get("columns") or [f["name"] for f in dv.schema_json["fields"]]
        limit = min(int(s.get("limit", 500)), 5000)
        rows = self.db.execute(select(Record).where(Record.dataset_version_id == dv.id).order_by(Record.row_index).limit(limit)).scalars()
        data = [[("[restricted]" if c in hidden else _cell(r.normalized.get(c))) for c in cols] for r in rows]
        return [{"kind": "table", "title": f"{ds.name} v{dv.version}", "columns": cols, "rows": data,
                 "caption": f"{len(data)} of {dv.record_count} records" + (f"; redacted: {', '.join(sorted(hidden & set(cols)))}" if hidden & set(cols) else "")}]

    def _s_analysis(self, s):
        run = self.db.get(AnalysisRun, s["analysis_run_id"])
        if run is None or run.project_id != self.project_id:
            raise NotFound("Analysis run referenced by report not found")
        for dsv in run.inputs.get("dataset_versions", []):
            self._dataset(dsv["dataset_id"], dsv["version"])
        self.manifest["analysis_runs"][run.id] = {"kind": run.kind, "params": run.params, "algorithm_version": run.algorithm_version,
                                                  "result_sha256": run.result_sha256, "created_at": str(run.created_at)}
        out = [{"kind": "paragraph", "text": f"Analysis {run.kind} (algorithm {run.algorithm_version}, run {run.id}, "
                                              f"result SHA-256 {run.result_sha256[:16]}…)."}]
        out.extend(analysis_blocks(run.kind, run.result))
        return out

    def _s_findings(self, s):
        inv = self._inv(s["investigation_id"])
        q = select(Finding).where(Finding.investigation_id == inv.id)
        if s.get("finding_ids"):
            q = q.where(Finding.id.in_(s["finding_ids"]))
        if not s.get("include_drafts"):
            q = q.where(Finding.status.in_(["confirmed_by_reviewer", "disputed", "under_review"]))
        out = []
        for f in self.db.execute(q.order_by(Finding.created_at)).scalars():
            links = self.db.execute(select(FindingEvidence, Evidence).join(Evidence, Evidence.id == FindingEvidence.evidence_id)
                                    .where(FindingEvidence.finding_id == f.id)).all()
            out.append({"kind": "finding", "title": f.title, "statement": f.statement, "status": f.status, "basis": f.basis,
                        "supporting": [_ev_ref(e) for l, e in links if l.stance == "supporting"],
                        "contradicting": [_ev_ref(e) for l, e in links if l.stance == "contradicting"]})
        if not out:
            out.append({"kind": "paragraph", "text": "No findings in reportable states."})
        return out

    def _s_evidence(self, s):
        inv = self._inv(s["investigation_id"])
        ev = self.db.execute(select(Evidence).where(Evidence.investigation_id == inv.id).order_by(Evidence.captured_at)).scalars().all()
        return [{"kind": "table", "title": "Evidence register", "columns": ["ID", "Source", "Location", "Excerpt", "Captured"],
                 "rows": [[e.id, f"{e.source_type}:{e.source_ref}", _loc(e), e.excerpt[:800], str(e.captured_at)[:19]] for e in ev]}]

    def _inv(self, inv_id: str) -> Investigation:
        inv = self.db.get(Investigation, inv_id)
        if inv is None or inv.project_id != self.project_id:
            raise NotFound("Investigation referenced by report not found")
        self.manifest["investigations"].append(inv.id)
        return inv

    def _s_graph(self, s):
        snap = self.db.get(GraphSnapshot, s["snapshot_id"])
        if snap is None or snap.project_id != self.project_id:
            raise NotFound("Graph snapshot not found")
        self.manifest["graph_snapshots"].append({"id": snap.id, "name": snap.name, "nodes": len(snap.node_ids), "edges": len(snap.edge_ids)})
        ents = {e.id: e for e in self.db.execute(select(Entity).where(Entity.id.in_(snap.node_ids))).scalars()}
        rels = self.db.execute(select(Relationship).where(Relationship.id.in_(snap.edge_ids))).scalars().all()
        nodes = [{"id": e.id, "label": e.display_name, "type": e.entity_type} for e in ents.values()]
        edges = [{"source": r.source_entity_id, "target": r.target_entity_id, "rel_type": r.rel_type, "assertion": r.assertion}
                 for r in rels if r.source_entity_id in ents and r.target_entity_id in ents]
        return [{"kind": "diagram", "title": snap.name, "nodes": nodes, "edges": edges},
                {"kind": "table", "title": "Relationships", "columns": ["Source", "Relationship", "Target", "Basis"],
                 "rows": [[ents[e["source"]].display_name, e["rel_type"], ents[e["target"]].display_name, e["assertion"]] for e in edges]}]

    def _s_timeline(self, s):
        q = select(Event).where(Event.project_id == self.project_id)
        if s.get("start"):
            q = q.where(Event.event_time >= datetime.fromisoformat(s["start"]))
        if s.get("end"):
            q = q.where(Event.event_time <= datetime.fromisoformat(s["end"]))
        if s.get("entity_ids"):
            q = q.where(Event.entity_id.in_(s["entity_ids"]))
        allowed = {d.id for d in self.db.execute(select(Dataset).where(Dataset.project_id == self.project_id)).scalars()
                   if d.access_policy != "restricted" or "dataset.read_restricted" in self.perms}
        evs = [e for e in self.db.execute(q.order_by(Event.event_time).limit(int(s.get("limit", 500)))).scalars()
               if e.dataset_id is None or e.dataset_id in allowed]
        by_month: dict[str, int] = {}
        for e in evs:
            k = e.event_time.strftime("%Y-%m")
            by_month[k] = by_month.get(k, 0) + 1
        return [{"kind": "chart", "chart": "bar", "title": "Events per month (UTC)", "labels": list(by_month), "values": list(by_month.values())},
                {"kind": "table", "title": "Events", "columns": ["Time (UTC)", "Title", "Category"],
                 "rows": [[e.event_time.strftime("%Y-%m-%d %H:%M"), e.title, e.category] for e in evs]}]

    def _s_map(self, s):
        q = select(GeoFeature).where(GeoFeature.project_id == self.project_id)
        if s.get("layers"):
            q = q.where(GeoFeature.layer.in_(s["layers"]))
        feats = self.db.execute(q.limit(5000)).scalars().all()
        pts = [{"lon": f.geometry["coordinates"][0], "lat": f.geometry["coordinates"][1], "label": f.name}
               for f in feats if f.geometry.get("type") == "Point"]
        return [{"kind": "scatter", "title": "Mapped points (WGS84)", "points": pts},
                {"kind": "table", "title": "Locations", "columns": ["Name", "Latitude", "Longitude", "Layer"],
                 "rows": [[f.name, f.geometry["coordinates"][1], f.geometry["coordinates"][0], f.layer] for f in feats
                          if f.geometry.get("type") == "Point"][:1000]}]


def _cell(v: Any) -> str:
    return "" if v is None else str(v)


def _loc(e: Evidence) -> str:
    l = e.location or {}
    if "page" in l:
        return f"{l.get('page_label', 'page')} {l['page']}"
    if "row_index" in l:
        return f"row {l['row_index'] + 1}"
    return ""


def _ev_ref(e: Evidence) -> str:
    return f"{e.id} ({e.source_type}{', ' + _loc(e) if _loc(e) else ''}): {e.excerpt[:200]}"


def analysis_blocks(kind: str, r: dict[str, Any]) -> list[dict[str, Any]]:
    if kind == "describe":
        cols = ["Field", "n", "Missing", "Distinct", "Mean", "Std", "Min", "Median", "Max"]
        return [{"kind": "table", "title": "Descriptive statistics", "columns": cols,
                 "rows": [[f["field"], f["n"], f["missing"], f["distinct"], *[_fmt(f.get(k)) for k in ("mean", "std", "min", "median", "max")]]
                          for f in r["fields"]]}]
    if kind == "frequency":
        return [{"kind": "chart", "chart": "bar", "title": f"Frequency of {r['field']}", "labels": [str(v['value']) for v in r["values"][:20]],
                 "values": [v["count"] for v in r["values"][:20]]},
                {"kind": "table", "title": "Frequency table", "columns": ["Value", "Count", "Share"],
                 "rows": [[v["value"], v["count"], _fmt(v["share"])] for v in r["values"]]}]
    if kind == "trend":
        return [{"kind": "chart", "chart": "line", "title": "Trend", "labels": [p["period"] for p in r["points"]],
                 "values": [p["value"] or 0 for p in r["points"]]},
                {"kind": "paragraph", "text": f"n={r['n_records']}; slope per period: {_fmt(r.get('linear_slope_per_period'))}. {r.get('assumptions', '')}"}]
    if kind == "outliers":
        return [{"kind": "paragraph", "text": f"{r.get('outlier_count', 0)} statistical outliers of n={r['n']}. {r.get('disclaimer', '')}"},
                {"kind": "table", "title": "Outliers", "columns": ["Row", "Value", "Robust z"],
                 "rows": [[o["row_index"] + 1, _fmt(o["value"]), _fmt(o["robust_z"])] for o in r.get("outliers", [])]}]
    if kind == "correlation":
        return [{"kind": "table", "title": "Correlations", "columns": ["A", "B", "n", "Pearson r", "95% CI"],
                 "rows": [[p["a"], p["b"], p["n"], _fmt(p["pearson_r"]), str(p["ci95"])] for p in r["pairs"]]},
                {"kind": "paragraph", "text": r.get("limitations", "")}]
    if "ranking" in r:
        return [{"kind": "table", "title": kind, "columns": ["Entity", "Score"], "rows": [[x["label"], _fmt(x["score"])] for x in r["ranking"]]}]
    return [{"kind": "code", "text": json.dumps(r, indent=2, default=str)[:20000]}]


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.4g}"
    return str(v)


# --------------------------------------------------------------------------- renderers


def render_html(title: str, blocks: list[dict[str, Any]], manifest: dict[str, Any]) -> bytes:
    e = html.escape
    parts = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(title)} — {APP_NAME}</title>",
             "<style>body{font-family:'DejaVu Sans',system-ui,sans-serif;max-width:960px;margin:2rem auto;padding:0 1rem;color:#1c2430}"
             "h1{border-bottom:3px solid #2f6f9f;padding-bottom:.3rem}h2{margin-top:2rem;color:#2f4f6f}"
             "table{border-collapse:collapse;width:100%;font-size:.85rem;margin:.5rem 0}th,td{border:1px solid #c9d2dc;padding:4px 6px;"
             "text-align:left;vertical-align:top;word-break:break-word}thead th{background:#e8eef4;position:sticky;top:0}"
             ".finding{border-left:4px solid #2f6f9f;padding:.5rem 1rem;background:#f5f8fb;margin:1rem 0}"
             ".badge{display:inline-block;padding:1px 6px;border-radius:3px;background:#dde6ee;font-size:.75rem}"
             "pre{background:#f3f5f7;padding:.5rem;overflow:auto}footer{margin-top:3rem;font-size:.75rem;color:#5a6672}</style></head><body>",
             f"<p class='badge'>{e(APP_NAME)} report</p><h1>{e(title)}</h1>",
             f"<p>Generated {e(manifest['generated_at'])} · Atlas {e(__version__)}</p>"]
    for b in blocks:
        k = b["kind"]
        if k == "heading":
            parts.append(f"<h2>{e(b['text'])}</h2>")
        elif k == "paragraph":
            parts.append(f"<p>{e(b['text'])}</p>")
        elif k == "code":
            parts.append(f"<pre>{e(b['text'])}</pre>")
        elif k == "table":
            parts.append(f"<h3>{e(b.get('title', ''))}</h3><table><thead><tr>" + "".join(f"<th>{e(str(c))}</th>" for c in b["columns"])
                         + "</tr></thead><tbody>" + "".join("<tr>" + "".join(f"<td>{e(_cell(v))}</td>" for v in row) + "</tr>" for row in b["rows"])
                         + "</tbody></table>" + (f"<p><small>{e(b['caption'])}</small></p>" if b.get("caption") else ""))
        elif k == "finding":
            parts.append(f"<div class='finding'><strong>{e(b['title'])}</strong> <span class='badge'>{e(b['status'])}</span> "
                         f"<span class='badge'>basis: {e(b['basis'])}</span><p>{e(b['statement'])}</p>"
                         + "<p><em>Supporting evidence</em></p><ul>" + "".join(f"<li>{e(x)}</li>" for x in b["supporting"]) + "</ul>"
                         + ("<p><em>Contradicting evidence</em></p><ul>" + "".join(f"<li>{e(x)}</li>" for x in b["contradicting"]) + "</ul>"
                            if b["contradicting"] else "") + "</div>")
        elif k == "chart":
            parts.append(svg_bar_chart(b))
        elif k == "diagram":
            parts.append(svg_diagram(b))
        elif k == "scatter":
            parts.append(svg_scatter(b))
    parts.append("<h2>Report manifest</h2><pre>" + e(json.dumps(manifest, indent=2, default=str)) + "</pre>")
    parts.append(f"<footer>{e(APP_NAME)} — Navigate the Complexity of Data.</footer></body></html>")
    return "".join(parts).encode("utf-8")


def svg_bar_chart(b: dict[str, Any]) -> str:
    vals = b["values"] or [0]
    w, h, pad = 800, 260, 40
    mx = max(vals) or 1
    n = max(len(vals), 1)
    bw = (w - 2 * pad) / n
    bars = []
    pts = []
    for i, v in enumerate(vals):
        x = pad + i * bw
        bh = (h - 2 * pad) * v / mx
        if b.get("chart") == "line":
            pts.append(f"{x + bw / 2:.1f},{h - pad - bh:.1f}")
        else:
            bars.append(f"<rect x='{x + 2:.1f}' y='{h - pad - bh:.1f}' width='{max(bw - 4, 1):.1f}' height='{bh:.1f}' fill='#2f6f9f'/>")
        if n <= 24:
            bars.append(f"<text x='{x + bw / 2:.1f}' y='{h - pad + 14}' font-size='9' text-anchor='middle'>{html.escape(str(b['labels'][i])[:12])}</text>")
    if pts:
        bars.append(f"<polyline fill='none' stroke='#2f6f9f' stroke-width='2' points='{' '.join(pts)}'/>")
    return (f"<h3>{html.escape(b.get('title', ''))}</h3><svg viewBox='0 0 {w} {h}' width='100%' role='img'>"
            f"<line x1='{pad}' y1='{h - pad}' x2='{w - pad}' y2='{h - pad}' stroke='#888'/>"
            f"<text x='{pad}' y='{pad - 8}' font-size='10'>max {mx}</text>{''.join(bars)}</svg>")


def _circle_layout(nodes: list[dict[str, Any]], w: float, h: float) -> dict[str, tuple[float, float]]:
    n = max(len(nodes), 1)
    r = min(w, h) / 2 - 60
    return {nd["id"]: (w / 2 + r * math.cos(2 * math.pi * i / n), h / 2 + r * math.sin(2 * math.pi * i / n))
            for i, nd in enumerate(nodes)}


def svg_diagram(b: dict[str, Any]) -> str:
    w, h = 800, 520
    pos = _circle_layout(b["nodes"], w, h)
    dash = {"observed": "", "inferred": "4,2", "model_suggested": "1,3", "manual": "6,3"}
    parts = [f"<line x1='{pos[e['source']][0]:.0f}' y1='{pos[e['source']][1]:.0f}' x2='{pos[e['target']][0]:.0f}' "
             f"y2='{pos[e['target']][1]:.0f}' stroke='#7a8a9a' stroke-dasharray='{dash.get(e['assertion'], '')}'/>" for e in b["edges"]]
    for nd in b["nodes"]:
        x, y = pos[nd["id"]]
        parts.append(f"<circle cx='{x:.0f}' cy='{y:.0f}' r='7' fill='#2f6f9f'/><text x='{x + 9:.0f}' y='{y + 4:.0f}' font-size='10'>"
                     f"{html.escape(nd['label'][:28])}</text>")
    return f"<h3>{html.escape(b['title'])}</h3><svg viewBox='0 0 {w} {h}' width='100%'>{''.join(parts)}</svg>" \
           "<p><small>Solid: observed · dashed: inferred · dotted: model-suggested · long dash: manual</small></p>"


def svg_scatter(b: dict[str, Any]) -> str:
    w, h = 800, 400
    pts = b["points"]
    if not pts:
        return "<p>No mapped points.</p>"
    lons, lats = [p["lon"] for p in pts], [p["lat"] for p in pts]
    x0, x1, y0, y1 = min(lons), max(lons), min(lats), max(lats)
    sx = (w - 40) / ((x1 - x0) or 1)
    sy = (h - 40) / ((y1 - y0) or 1)
    s = min(sx, sy)
    circles = "".join(f"<circle cx='{20 + (p['lon'] - x0) * s:.1f}' cy='{h - 20 - (p['lat'] - y0) * s:.1f}' r='3' fill='#c0392b'/>" for p in pts)
    return (f"<h3>{html.escape(b['title'])}</h3><svg viewBox='0 0 {w} {h}' width='100%' style='background:#eef2f5'>{circles}</svg>"
            f"<p><small>Equirectangular plot, lon {x0:.3f}…{x1:.3f}, lat {y0:.3f}…{y1:.3f}. No basemap (offline).</small></p>")


def render_pdf(title: str, blocks: list[dict[str, Any]], manifest: dict[str, Any]) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Circle, Drawing, Line, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import KeepTogether, LongTable, Paragraph, Preformatted, SimpleDocTemplate, Spacer, TableStyle

    if "AtlasSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AtlasSans", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("AtlasSans-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        from reportlab.pdfbase.pdfmetrics import registerFontFamily

        registerFontFamily("AtlasSans", normal="AtlasSans", bold="AtlasSans-Bold", italic="AtlasSans", boldItalic="AtlasSans-Bold")
    base = ParagraphStyle("base", fontName="AtlasSans", fontSize=9.5, leading=13)
    h1 = ParagraphStyle("h1", parent=base, fontName="AtlasSans-Bold", fontSize=18, leading=22, spaceAfter=6,
                        textColor=colors.HexColor("#1f3b57"))
    h2 = ParagraphStyle("h2", parent=base, fontName="AtlasSans-Bold", fontSize=13, leading=16, spaceBefore=10, spaceAfter=4,
                        textColor=colors.HexColor("#2f4f6f"))
    h3 = ParagraphStyle("h3", parent=base, fontName="AtlasSans-Bold", fontSize=10.5, spaceBefore=6, spaceAfter=3)
    cell = ParagraphStyle("cell", parent=base, fontSize=7.5, leading=9.5, wordWrap="CJK")
    cellh = ParagraphStyle("cellh", parent=cell, fontName="AtlasSans-Bold")
    mono = ParagraphStyle("mono", parent=base, fontName="AtlasSans", fontSize=6.5, leading=8)
    page_w, page_h = A4
    margin = 18 * mm
    avail = page_w - 2 * margin

    def esc(t: Any) -> str:
        return html.escape(_cell(t)).replace("\n", "<br/>")

    story: list[Any] = [Paragraph(esc(title), h1),
                        Paragraph(esc(f"{APP_NAME} · generated {manifest['generated_at']} · software {__version__}"), base), Spacer(1, 6)]
    for b in blocks:
        k = b["kind"]
        if k == "heading":
            story.append(Paragraph(esc(b["text"]), h2))
        elif k == "paragraph":
            story.append(Paragraph(esc(b["text"]), base))
        elif k == "code":
            story.append(Preformatted(b["text"][:15000], mono))
        elif k == "table":
            cols = [str(c) for c in b["columns"]]
            if b.get("title"):
                story.append(Paragraph(esc(b["title"]), h3))
            rows = b["rows"] or [["(no rows)"] + [""] * (len(cols) - 1)]
            # column widths proportional to content length, bounded, summing to available width
            lens = [max([len(c)] + [min(len(_cell(r[i])), 80) for r in rows[:200]]) for i, c in enumerate(cols)]
            lens = [min(max(x, 4), 60) for x in lens]
            total = sum(lens)
            widths = [avail * x / total for x in lens]
            data = [[Paragraph(esc(c), cellh) for c in cols]] + [[Paragraph(esc(v)[:3000], cell) for v in r] for r in rows]
            t = LongTable(data, colWidths=widths, repeatRows=1, splitByRow=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c4d0")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e3eaf1")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fa")])]))
            story.append(t)
            if b.get("caption"):
                story.append(Paragraph(esc(b["caption"]), cell))
        elif k == "finding":
            items = [Paragraph(f"<b>{esc(b['title'])}</b> — status: {esc(b['status'])}; basis: {esc(b['basis'])}", base),
                     Paragraph(esc(b["statement"]), base), Paragraph("<i>Supporting evidence</i>", base)]
            items += [Paragraph("• " + esc(x), cell) for x in b["supporting"]] or [Paragraph("(none)", cell)]
            if b["contradicting"]:
                items.append(Paragraph("<i>Contradicting evidence</i>", base))
                items += [Paragraph("• " + esc(x), cell) for x in b["contradicting"]]
            story.append(KeepTogether(items + [Spacer(1, 6)]))
        elif k == "chart":
            vals = [float(v) for v in b["values"]] or [0.0]
            d = Drawing(avail, 180)
            ch = HorizontalLineChart() if b.get("chart") == "line" else VerticalBarChart()
            ch.x, ch.y, ch.width, ch.height = 30, 30, avail - 50, 130
            ch.data = [vals]
            ch.categoryAxis.categoryNames = [str(l)[:10] for l in b["labels"]] or [""]
            ch.categoryAxis.labels.fontName = "AtlasSans"
            ch.categoryAxis.labels.fontSize = 6
            ch.categoryAxis.labels.angle = 45 if len(vals) > 8 else 0
            ch.categoryAxis.labels.boxAnchor = "ne" if len(vals) > 8 else "n"
            ch.valueAxis.labels.fontName = "AtlasSans"
            ch.valueAxis.labels.fontSize = 7
            ch.valueAxis.valueMin = 0
            if b.get("chart") == "line":
                ch.lines[0].strokeColor = colors.HexColor("#2f6f9f")
            else:
                ch.bars[0].fillColor = colors.HexColor("#2f6f9f")
            d.add(ch)
            d.add(String(0, 168, b.get("title", ""), fontName="AtlasSans-Bold", fontSize=9))
            story.append(d)
        elif k == "diagram":
            dh = 320
            d = Drawing(avail, dh)
            pos = _circle_layout(b["nodes"], avail, dh)
            for e in b["edges"]:
                (x1, y1), (x2, y2) = pos[e["source"]], pos[e["target"]]
                ln = Line(x1, dh - y1, x2, dh - y2, strokeColor=colors.HexColor("#7a8a9a"), strokeWidth=0.7)
                if e["assertion"] != "observed":
                    ln.strokeDashArray = [3, 2]
                d.add(ln)
            for nd in b["nodes"]:
                x, y = pos[nd["id"]]
                d.add(Circle(x, dh - y, 4, fillColor=colors.HexColor("#2f6f9f"), strokeColor=None))
                d.add(String(x + 6, dh - y - 3, nd["label"][:28], fontName="AtlasSans", fontSize=6.5))
            story += [Paragraph(esc(b["title"]), h3), d,
                      Paragraph("Solid lines: observed relationships; dashed: inferred, model-suggested or manual.", cell)]
        elif k == "scatter":
            pts = b["points"]
            if not pts:
                story.append(Paragraph("No mapped points.", base))
                continue
            dh = 220
            d = Drawing(avail, dh)
            lons, lats = [p["lon"] for p in pts], [p["lat"] for p in pts]
            x0, x1, y0, y1 = min(lons), max(lons), min(lats), max(lats)
            s = min((avail - 20) / ((x1 - x0) or 1), (dh - 20) / ((y1 - y0) or 1))
            for p in pts:
                d.add(Circle(10 + (p["lon"] - x0) * s, 10 + (p["lat"] - y0) * s, 2, fillColor=colors.HexColor("#c0392b"), strokeColor=None))
            story += [Paragraph(esc(b["title"]), h3), d,
                      Paragraph(esc(f"Equirectangular plot, lon {x0:.3f}…{x1:.3f}, lat {y0:.3f}…{y1:.3f}; no basemap."), cell)]
    story += [Paragraph("Report manifest", h2), Preformatted(json.dumps(manifest, indent=1, default=str)[:20000], mono)]

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("AtlasSans", 7)
        canvas.setFillColor(colors.HexColor("#5a6672"))
        canvas.drawString(margin, 10 * mm, f"{APP_NAME} — {title[:80]}")
        canvas.drawRightString(page_w - margin, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=18 * mm,
                            title=title, author=APP_NAME, creator=f"{APP_NAME} {__version__}")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def render_tables_csv(blocks: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for b in blocks:
        if b["kind"] == "table":
            w.writerow([f"# {b.get('title', '')}"])
            w.writerow(b["columns"])
            for r in b["rows"]:
                w.writerow([_csv_safe(_cell(v)) for v in r])
            w.writerow([])
    return buf.getvalue().encode("utf-8-sig")


def _csv_safe(v: str) -> str:
    # Prevent spreadsheet formula injection when exports are opened in Excel-like tools.
    return "'" + v if v[:1] in {"=", "+", "-", "@", "\t", "\r"} and not _is_number(v) else v


def _is_number(v: str) -> bool:
    try:
        float(v)
        return True
    except ValueError:
        return False


def render_xlsx(title: str, blocks: list[dict[str, Any]], manifest: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Manifest"
    ws0.append([title])
    ws0["A1"].font = Font(bold=True, size=14)
    for k, v in manifest.items():
        ws0.append([k, json.dumps(v, default=str)])
    n = 0
    for b in blocks:
        if b["kind"] != "table":
            continue
        n += 1
        ws = wb.create_sheet(f"T{n} {(b.get('title') or '')[:24]}".replace("/", "-").replace(":", "-").replace("[", "(").replace("]", ")")
                             .replace("*", "").replace("?", "").replace("\\", "-"))
        ws.append([str(c) for c in b["columns"]])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in b["rows"]:
            ws.append([_csv_safe(_cell(v)) for v in r])
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_report(ctx: "AtlasContext", db: Session, *, project_id: str, report_id: str, fmt: str, user_id: str,
                  perms: set[str]) -> ReportExport:
    if fmt not in FORMATS:
        raise ValidationFailed(f"Format must be one of {sorted(FORMATS)}")
    rpt = db.get(Report, report_id)
    if rpt is None or rpt.project_id != project_id:
        raise NotFound("Report not found")
    validate_definition(rpt.definition)
    res = Resolver(db, project_id, perms)
    blocks = res.resolve(rpt.definition)
    generated = datetime.now(timezone.utc).isoformat()
    manifest = {"report_id": rpt.id, "title": rpt.title, "format": fmt, "generated_at": generated, "generated_by": user_id,
                "software": {"atlas": __version__, "renderer": RENDERER_VERSION, "python": platform.python_version()},
                **res.manifest, "definition": rpt.definition}
    if fmt == "pdf":
        data = render_pdf(rpt.title, blocks, manifest)
    elif fmt == "html":
        data = render_html(rpt.title, blocks, manifest)
    elif fmt == "json":
        data = json.dumps({"title": rpt.title, "blocks": blocks, "manifest": manifest}, indent=2, default=str).encode()
    elif fmt == "csv":
        data = render_tables_csv(blocks)
    else:
        data = render_xlsx(rpt.title, blocks, manifest)
    sha = hashlib.sha256(data).hexdigest()
    d = ctx.settings.exports_dir / project_id / "reports"
    d.mkdir(parents=True, exist_ok=True)
    stamp = generated[:19].replace(":", "").replace("-", "")
    path = d / f"{rpt.id}-{stamp}.{fmt}"
    path.write_bytes(data)
    manifest["output_sha256"] = sha
    (path.with_suffix(path.suffix + ".manifest.json")).write_text(json.dumps(manifest, indent=2, default=str), encoding="utf-8")
    rex = ReportExport(report_id=rpt.id, format=fmt, stored_path=str(path.relative_to(ctx.settings.storage_dir)), sha256=sha,
                       size_bytes=len(data), manifest=json.loads(json.dumps(manifest, default=str)), created_by=user_id)
    db.add(rex)
    db.flush()
    return rex
